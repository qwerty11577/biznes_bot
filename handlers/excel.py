from aiogram import Router, F
from aiogram.types import Message, BufferedInputFile
from aiogram.filters import Command
import database as db
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import io
from datetime import datetime

router = Router()

@router.message(F.text == "📥 Excel отчёт")
async def excel_report(message: Message):
    user_id = message.from_user.id
    
    wb = openpyxl.Workbook()
    
    # Транзакции
    ws1 = wb.active
    ws1.title = "Доходы и расходы"
    ws1.append(["Тип", "Сумма", "Категория", "Описание", "Дата"])
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    transactions = await db.get_transactions(user_id, limit=1000)
    for t in transactions:
        type_, amount, category, desc, date = t
        ws1.append([
            "Доход" if type_ == "daromad" else "Расход",
            amount, category, desc, date
        ])
    
    # Долги
    ws2 = wb.create_sheet("Долги")
    ws2.append(["Человек", "Сумма", "Тип", "Описание"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    debts = await db.get_debts(user_id)
    for d in debts:
        id_, person, amount, type_, desc = d
        ws2.append([
            person, amount,
            "Мне должны" if type_ == "menga" else "Я должен",
            desc
        ])
    
    # Товары
    ws3 = wb.create_sheet("Товары")
    ws3.append(["Название", "Количество", "Цена", "Единица", "Итого"])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    products = await db.get_products(user_id)
    for p in products:
        id_, name, qty, price, unit = p
        ws3.append([name, qty, price, unit, qty * price])
    
    # Сохраняем в память
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"hisobot_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    file = BufferedInputFile(buffer.read(), filename=filename)
    
    await message.answer_document(file, caption="📊 Ваш отчёт готов!")